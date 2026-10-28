"""Build a Koach Hitarvut letter (.docx) matching the official template.

Template measurements (from the reference PDF, A4 595x842pt):
  logo 125x105pt centered in header, 9pt from top
  body David 12pt, line pitch 17.7pt, paragraph gap ~10pt
  margins: right 50pt, left 45pt, body starts ~143pt from top
Usage: python3 build_letter.py content.json logo.png out.docx
"""
import json
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "David"
SIZE = 12
PITCH = 17.7
# Defaults reproduce the template; content["layout"] may tighten them to fit one page
LAYOUT = {"size": SIZE, "pitch": PITCH, "gap": 1.0, "top": 140}


def set_rtl(par, jc=None):
    ppr = par._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.insert(0, bidi)
    if jc:
        el = OxmlElement("w:jc")
        el.set(qn("w:val"), jc)
        ppr.append(el)


def fmt(par, before=0, after=0, pitch=None, right_indent=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before * LAYOUT["gap"])
    pf.space_after = Pt(after * LAYOUT["gap"])
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = pitch or Pt(LAYOUT["pitch"])
    if right_indent is not None:
        # In a bidi paragraph w:left is the visual right (start) side
        ppr = par._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(int(right_indent * 20)))
        ppr.append(ind)


def run(par, text, bold=False, underline=False, size=None):
    size = size or LAYOUT["size"]
    r = par.add_run(text)
    rpr = r._r.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    rpr.insert(0, fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
        rpr.append(OxmlElement("w:bCs"))
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(int(size * 2)))
        rpr.append(el)
    rpr.append(OxmlElement("w:rtl"))
    return r


def para(doc, jc=None, **kw):
    p = doc.add_paragraph()
    set_rtl(p, jc)
    fmt(p, **kw)
    return p


def build(content, logo, out):
    LAYOUT.update(content.get("layout", {}))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(595.3), Pt(841.9)
    sec.right_margin, sec.left_margin = Pt(50), Pt(45)
    sec.top_margin, sec.bottom_margin = Pt(LAYOUT["top"]), Pt(42)
    sec.header_distance, sec.footer_distance = Pt(9), Pt(30)
    sec._sectPr.append(OxmlElement("w:bidi"))

    hp = sec.header.paragraphs[0]
    set_rtl(hp, "center")
    hp.add_run().add_picture(logo, width=Pt(125.4), height=Pt(105.2))

    # Date block: visually left-aligned (jc=right in a bidi paragraph)
    para(doc, after=0, pitch=Pt(11))
    run(para(doc, jc="right", after=2.3), content["date_hebrew"])
    run(para(doc, jc="right", after=11.3), content["date_gregorian"])

    run(para(doc, after=0.3), "לכבוד:", bold=True)
    for i, line in enumerate(content["to"]):
        run(para(doc, after=3.3 if i < len(content["to"]) - 1 else 0), line)
    para(doc, after=0)

    p = para(doc, jc="center", before=0, after=6.3)
    run(p, "הנדון: ", bold=True, underline=True)
    run(p, content["subject"], bold=True, underline=True)
    if content.get("subtitle"):
        run(para(doc, jc="center", after=9.3), content["subtitle"], bold=True)

    n = 0
    for sec_ in content["sections"]:
        if sec_.get("title"):  # untitled sections give a plain numbered list
            run(para(doc, before=0, after=8.3), sec_["title"], bold=True, underline=True)
        for item in sec_["paragraphs"]:
            if isinstance(item, dict):  # lettered sub-list
                for j, sub in enumerate(item["items"]):
                    last = j == len(item["items"]) - 1
                    p = para(doc, jc="both", after=10.3 if last else 6.3, right_indent=36)
                    run(p, sub)
            else:
                n += 1
                p = para(doc, jc="both", after=10.3)
                run(p, f"{n}. {item}")

    run(para(doc, jc="center", before=0, after=0.3), "בברכה,", bold=True)
    run(para(doc, jc="center", after=6.3), content["signer"])
    run(para(doc, jc="center", after=0), content["org"])

    doc.save(out)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2], sys.argv[3])
